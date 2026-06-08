"""
Generate a TraceRoot project presentation (.pptx) from the current README and app.py contents.

Run this file after you update project documentation or code to refresh the presentation.
"""

import os
from datetime import datetime

try:
    from pptx import Presentation
    from pptx.util import Pt
except ImportError as exc:
    raise ImportError(
        "python-pptx is required to run this script. Install it with: pip install python-pptx"
    ) from exc

try:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, ListFlowable, ListItem
    from reportlab.lib.units import inch
    from reportlab.lib import colors
except ImportError as exc:
    raise ImportError(
        "reportlab is required to generate the PDF. Install it with: pip install reportlab"
    ) from exc

README_PATH = os.path.join(os.path.dirname(__file__), 'README.md')
APP_PATH = os.path.join(os.path.dirname(__file__), 'app.py')
OUTPUT_FILE = os.path.join(os.path.dirname(__file__), 'TraceRoot_Project_Presentation.pptx')


def read_markdown_sections(path):
    """Parse README headings and return title, summary, and section content."""
    with open(path, encoding='utf-8') as f:
        lines = f.read().splitlines()

    title = ''
    summary_lines = []
    sections = {}
    current_section = None
    current_lines = []
    in_code = False

    for line in lines:
        stripped = line.strip()
        if stripped.startswith('```'):
            in_code = not in_code
            continue
        if in_code:
            continue

        if stripped.startswith('# '):
            title = stripped[2:].strip()
            continue

        if stripped.startswith('## '):
            if current_section:
                sections[current_section] = current_lines
            current_section = stripped[3:].strip()
            current_lines = []
            continue

        if current_section is None:
            if stripped:
                summary_lines.append(stripped)
        else:
            current_lines.append(line)

    if current_section:
        sections[current_section] = current_lines

    summary = ' '.join([line for line in summary_lines if line])
    return title, summary, sections


def extract_routes(path):
    """Extract Flask route patterns from app.py."""
    routes = []
    with open(path, encoding='utf-8') as f:
        for line in f:
            if '@app.route' in line:
                route_part = line.split('(', 1)[1].split(')', 1)[0]
                route_part = route_part.strip().strip('"\'"')
                if route_part:
                    routes.append(route_part)
    return routes


def bullets(lines):
    """Convert markdown section lines into bullet list items."""
    items = []
    for raw_line in lines:
        line = raw_line.strip()
        if line.startswith('- '):
            items.append(line[2:].strip())
        elif line.startswith('* '):
            items.append(line[2:].strip())
    return items


def paragraph(lines, max_sentences=3):
    joined = ' '.join([line.strip() for line in lines if line.strip() and not line.strip().startswith('-')])
    if not joined:
        return ''
    sentences = joined.split('.')
    condensed = '. '.join([s.strip() for s in sentences if s.strip()][:max_sentences])
    if condensed and not condensed.endswith('.'):
        condensed += '.'
    return condensed


def add_bullet_slide(prs, title, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    body = slide.shapes.placeholders[1].text_frame
    body.clear()
    body.word_wrap = True
    for bullet in bullets:
        paragraph = body.add_paragraph()
        paragraph.text = bullet
        paragraph.level = 0
        paragraph.font.size = Pt(18)
    return slide


def create_presentation(title, summary, sections, routes):
    prs = Presentation()

    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = title
    subtitle = title_slide.placeholders[1]
    subtitle.text = f"Auto-generated presentation for TraceRoot Local\nLast refreshed: {datetime.now():%Y-%m-%d %H:%M}" 

    problem_bullets = [
        "Need a lightweight local bug analysis platform for fast team review.",
        "Manual bug RCA is slow and hard to maintain for demo or evaluation.",
        "A self-contained system keeps the workflow reproducible and easy to run.",
    ]
    add_bullet_slide(prs, "Problem Statement", problem_bullets)

    initiative_bullets = [
        "Deliver a simple Flask-based platform for bug analysis and RCA.",
        "Use synthetic bug data, month filtering, and search to simulate real workflows.",
        "Integrate a local Ollama model for Root Cause Analysis generation.",
        "Keep deployment lightweight with no external database required.",
    ]
    add_bullet_slide(prs, "Project Initiative", initiative_bullets)

    overview_bullets = []
    if summary:
        overview_bullets.append(summary)
    if not overview_bullets:
        overview_bullets.append("TraceRoot Local is a Flask app for bug reporting, analysis, and RCA.")
    add_bullet_slide(prs, "Project Overview", overview_bullets)

    tech_bullets = bullets(sections.get('Technical Stack', []))
    if not tech_bullets:
        tech_bullets = bullets(sections.get('🎯 Features', []))[:4]
    add_bullet_slide(prs, "Technical Stack", tech_bullets)

    feature_bullets = bullets(sections.get('🎯 Features', []))
    if feature_bullets:
        add_bullet_slide(prs, "Key Features", feature_bullets[:6])

    if routes:
        route_bullets = [f"{route}" for route in routes]
        add_bullet_slide(prs, "Flask Routes", route_bullets)

    data_model_bullets = bullets(sections.get('📊 Data Structure', []))
    if data_model_bullets:
        add_bullet_slide(prs, "Data Model", data_model_bullets[:8])

    ai_bullets = bullets(sections.get('🤖 RCA Generation (Ollama Integration)', []))
    if ai_bullets:
        add_bullet_slide(prs, "AI Integration", ai_bullets)

    usage_lines = sections.get('🎮 Usage', [])
    usage_bullets = []
    for line in usage_lines:
        stripped = line.strip()
        if stripped.startswith('- ') or stripped.startswith('1.') or stripped.startswith('2.') or stripped.startswith('3.'):
            usage_bullets.append(stripped.lstrip(' -').strip())
    if usage_bullets:
        add_bullet_slide(prs, "Usage", usage_bullets[:8])

    add_bullet_slide(prs, "Refresh Instructions", [
        "Edit the code or README content in the repository.",
        "Run: python ppt.py",
        f"The presentation file will be updated at: {os.path.basename(OUTPUT_FILE)}",
        "Make sure python-pptx is installed before running the script.",
    ])

    prs.save(OUTPUT_FILE)
    return OUTPUT_FILE


def extract_text_from_presentation(prs):
    """Extract text from slides to reuse in the PDF output."""
    slides = []
    for slide in prs.slides:
        title = ''
        bullets = []
        for shape in slide.shapes:
            if not hasattr(shape, 'has_text_frame') or not shape.has_text_frame:
                continue
            text = shape.text.strip()
            if not text:
                continue

            if shape == slide.shapes.title:
                title = text
            else:
                for line in text.splitlines():
                    line = line.strip()
                    if line:
                        bullets.append(line)

        slides.append((title or 'Slide', bullets))
    return slides


def generate_pdf_from_presentation(prs, pdf_path):
    """Create a PDF file from the generated presentation content."""
    slides = extract_text_from_presentation(prs)
    doc = SimpleDocTemplate(pdf_path, pagesize=letter,
                            rightMargin=72, leftMargin=72,
                            topMargin=72, bottomMargin=72)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'Title', parent=styles['Heading1'], fontSize=28,
        textColor=colors.HexColor('#0f172a'), spaceAfter=12, leading=32
    )
    heading_style = ParagraphStyle(
        'Heading', parent=styles['Heading2'], fontSize=18,
        textColor=colors.HexColor('#1d4ed8'), spaceAfter=10, leading=22
    )
    bullet_style = ParagraphStyle(
        'Bullet', parent=styles['BodyText'], fontSize=12,
        leading=18, leftIndent=18, firstLineIndent=-12, spaceAfter=4,
        bulletIndent=0, textColor=colors.HexColor('#111111')
    )
    body_style = ParagraphStyle(
        'Body', parent=styles['BodyText'], fontSize=12,
        leading=18, textColor=colors.HexColor('#111111'), spaceAfter=12
    )

    story = []
    for idx, (slide_title, bullets) in enumerate(slides):
        story.append(Paragraph(slide_title, heading_style))
        story.append(Spacer(1, 0.08 * inch))

        if bullets:
            list_items = []
            for bullet in bullets:
                list_items.append(ListItem(Paragraph(bullet, bullet_style), bulletColor=colors.HexColor('#0f172a')))
            story.append(ListFlowable(list_items, bulletType='bullet', start='•', leftIndent=12, bulletFontName='Helvetica', bulletFontSize=12))
        else:
            story.append(Paragraph('No content available for this slide.', body_style))

        if idx < len(slides) - 1:
            story.append(PageBreak())

    doc.build(story)
    return pdf_path


def main():
    title, summary, sections = read_markdown_sections(README_PATH)
    routes = extract_routes(APP_PATH)
    output_path = create_presentation(title or 'TraceRoot Local', summary, sections, routes)
    print(f"Presentation generated: {output_path}")

    pdf_path = os.path.splitext(output_path)[0] + '.pdf'
    prs = Presentation(output_path)
    pdf_output = generate_pdf_from_presentation(prs, pdf_path)
    print(f"PDF generated: {pdf_output}")


if __name__ == '__main__':
    main()
